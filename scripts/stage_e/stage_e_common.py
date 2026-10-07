from __future__ import annotations
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any
import json, re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PERIODS={"Baseline","Term 1","Term 2","Term 3"}
MONTHS={m:i for i,m in enumerate(["January","February","March","April","May","June","July","August","September","October","November","December"],1)}
ATTN_GRADE={"LKG_Attn":"LKG","UKG_Attn":"UKG",**{f"Class{i}_Attn":f"Class {i}" for i in range(1,9)}}
CANONICAL_GRADES=["LKG","UKG"]+[f"Class {i}" for i in range(1,9)]
GRADE_LABEL_MAP={"A. LKG":"LKG","B. UKG":"UKG"}
ATTENDANCE_MEANINGS={"P":"Present","A":"Absent","H":"Holiday","NA":"Not Applicable"}
ASSESSMENT_SCORE_MEANINGS={"A":"Absent"}
FAILURE_SHEET_ORDER=["Students & Attendance","UKG","LKG","Class 8","Class 7","Class 6","Class 5","Class 4","Class 3","Class 2","Class 1"]

def blank(v): return v is None or (isinstance(v,str) and v.strip()=="")
def sval(v): return None if blank(v) else str(v).strip()
def numeric(v):
    if blank(v): return None
    try: return float(v)
    except (TypeError,ValueError): return None
def as_int(v):
    n=numeric(v)
    return int(n) if n is not None and n.is_integer() else None
def as_date(v):
    if blank(v): return None
    if isinstance(v,datetime): return v.date()
    if isinstance(v,date): return v
    for fmt in ("%d-%b-%Y","%Y-%m-%d","%d/%m/%Y"):
        try: return datetime.strptime(str(v).strip(),fmt).date()
        except ValueError: pass
    return None
def rows_values(ws): return list(ws.iter_rows(values_only=True))
def load_manifest(path): return json.loads(Path(path).read_text(encoding="utf-8"))
def remember(d,key,**prov): d.setdefault(key,prov)

def parse_assessment_score(raw):
    value=sval(raw)
    if value=="A":
        return None,"A"
    try:
        n=float(value)
    except (TypeError,ValueError):
        return None
    if not n.is_integer() or not 0<=n<=10:
        return None
    return int(n),None

def assessment_effective_score(numeric_score,score_code):
    if numeric_score is not None and score_code is None:
        if 0<=numeric_score<=10:
            return numeric_score
        raise ValueError("numeric assessment score must be between 0 and 10")
    if numeric_score is None and score_code=="A":
        return 0
    raise ValueError("unsupported or ambiguous assessment score representation")

def calculate_assessment_average(score_rows):
    rows=list(score_rows)
    if not rows:
        raise ValueError("assessment average is undefined without score observations")
    effective=[assessment_effective_score(r.get("numeric_score"),r.get("score_code")) for r in rows]
    return sum(effective)/len(effective)

@dataclass
class StageEData:
    manifest: dict
    source_files:list=field(default_factory=list)
    academic_years:dict=field(default_factory=dict)
    grades:dict=field(default_factory=dict)
    subjects:dict=field(default_factory=dict)
    periods:dict=field(default_factory=dict)
    score_codes:dict=field(default_factory=dict)
    attendance_codes:dict=field(default_factory=dict)
    students:list=field(default_factory=list)
    enrollments:list=field(default_factory=list)
    grade_subjects:list=field(default_factory=list)
    competencies:list=field(default_factory=list)
    activities:list=field(default_factory=list)
    assessments:list=field(default_factory=list)
    scores:list=field(default_factory=list)
    attendance_months:list=field(default_factory=list)
    attendance_days:list=field(default_factory=list)
    measurements:list=field(default_factory=list)
    failures:dict=field(default_factory=lambda:defaultdict(list))
    summary:dict=field(default_factory=dict)

def parse_sources(source_dir:Path,manifest_path:Path)->StageEData:
    manifest=load_manifest(manifest_path); data=StageEData(manifest)
    by_label={w["label"]:w for w in manifest["workbooks"]}
    for w in manifest["workbooks"]:
        p=source_dir/w["filename"]
        if not p.exists(): raise FileNotFoundError(f"Missing expected sanitized workbook: {p}")
        data.source_files.append({**w,"path":p})

    wb=load_workbook(source_dir/by_label["Students & Attendance"]["filename"],data_only=True,read_only=True)
    student_rows=rows_values(wb["Students"]); headers=list(student_rows[0])
    valid_students=set(); student_row_map={}; student_source_rows=student_failed_rows=0
    for c in (8,9,10):
        m=re.search(r"(\d{4}-\d{2})",sval(headers[c-1]) or "")
        if m: remember(data.academic_years,m.group(1),source_label="Students & Attendance",sheet="Students",row=1)
    for idx,vals0 in enumerate(student_rows[1:],2):
        vals=list(vals0[:10])+[None]*max(0,10-len(vals0))
        if all(blank(v) for v in vals): continue
        student_source_rows+=1; sid,name=sval(vals[0]),sval(vals[1]); bad=False
        if not sid:
            data.failures["Students & Attendance"].append(["Students",idx,"A","student","Missing required student_id",vals[0]]); bad=True
        if not name:
            data.failures["Students & Attendance"].append(["Students",idx,"B","student","Missing required student_name",vals[1]]); bad=True
        if bad: student_failed_rows+=1; continue
        valid_students.add(sid); student_row_map[sid]=(idx,vals)
        data.students.append({"student_id":sid,"student_name":name,"source_enrollment_year_label":sval(vals[2]),"gender":sval(vals[3]),"date_of_birth":as_date(vals[4]),"social_category":sval(vals[5]),"program_type":sval(vals[6]),"source_label":"Students & Attendance","sheet":"Students","row":idx})
        if sval(vals[2]): remember(data.academic_years,sval(vals[2]),source_label="Students & Attendance",sheet="Students",row=idx)
    for order,grade in enumerate(CANONICAL_GRADES):
        sheet="LKG_Attn" if grade=="LKG" else "UKG_Attn" if grade=="UKG" else grade.replace(" ","")+"_Attn"
        data.grades[grade]={"grade_code":grade,"grade_name":grade,"sort_order":order,"source_label":"Students & Attendance","sheet":sheet,"row":1}

    enrollment_source_cells=enrollment_failures=0; unresolved=Counter()
    for sid,(r,vals) in student_row_map.items():
        for c in (8,9,10):
            raw=vals[c-1]
            if blank(raw): continue
            enrollment_source_cells+=1; header=sval(headers[c-1]) or ""; m=re.search(r"(\d{4}-\d{2})",header); label=sval(raw)
            if not m:
                data.failures["Students & Attendance"].append(["Students",r,get_column_letter(c),"enrollment","Cannot derive academic year from class column header",header]); enrollment_failures+=1; continue
            year=m.group(1); remember(data.academic_years,year,source_label="Students & Attendance",sheet="Students",row=1)
            if label in GRADE_LABEL_MAP: canonical=GRADE_LABEL_MAP[label]
            elif label in CANONICAL_GRADES: canonical=label
            else:
                data.failures["Students & Attendance"].append(["Students",r,get_column_letter(c),"enrollment","Unrecognized source grade label",label]); enrollment_failures+=1; continue
            data.enrollments.append({"student_id":sid,"academic_year":year,"grade":canonical,"source_grade_label":label,"source_label":"Students & Attendance","sheet":"Students","row":r,"column":get_column_letter(c)})

    attendance_source_rows=attendance_failed_rows=0; attendance_status_norm=Counter(); transforms=Counter(); attendance_keys=[]
    for sheet,grade in ATTN_GRADE.items():
        for r,vals0 in enumerate(rows_values(wb[sheet])[1:],2):
            vals=list(vals0[:43])+[None]*max(0,43-len(vals0))
            if all(blank(v) for v in vals): continue
            attendance_source_rows+=1; sid,ay,month=sval(vals[0]),sval(vals[2]),sval(vals[3]); bad=False
            if not sid:
                data.failures["Students & Attendance"].append([sheet,r,"A","attendance_month","Missing required student_id",vals[0]]); bad=True
            elif sid not in valid_students:
                data.failures["Students & Attendance"].append([sheet,r,"A","attendance_month","student_id not present in valid Students master row",sid]); bad=True
            if not ay:
                data.failures["Students & Attendance"].append([sheet,r,"C","attendance_month","Missing required academic_year",vals[2]]); bad=True
            if not month:
                data.failures["Students & Attendance"].append([sheet,r,"D","attendance_month","Missing required month",vals[3]]); bad=True
            elif month not in MONTHS:
                data.failures["Students & Attendance"].append([sheet,r,"D","attendance_month","Unrecognized month label",month]); bad=True
            if bad: attendance_failed_rows+=1; continue
            remember(data.academic_years,ay,source_label="Students & Attendance",sheet=sheet,row=r)
            attendance_keys.append((sid,ay,grade,MONTHS[month]))
            data.attendance_months.append({"student_id":sid,"academic_year":ay,"grade":grade,"month_number":MONTHS[month],"source_month_label":month,"reported_days_present":as_int(vals[40]),"reported_days_absent":as_int(vals[41]),"reported_attendance_percentage":numeric(vals[42]),"source_label":"Students & Attendance","sheet":sheet,"row":r})
            for c in range(10,41):
                raw=vals[c-1]
                if blank(raw): continue
                raw_s=str(raw); norm=raw_s.strip()
                if not norm:
                    data.failures["Students & Attendance"].append([sheet,r,get_column_letter(c),"attendance_day","Attendance code is whitespace-only",raw_s]); continue
                attendance_status_norm[norm]+=1
                if raw_s!=norm: transforms[(raw_s,norm)]+=1
                remember(data.attendance_codes,norm,source_label="Students & Attendance",sheet=sheet,row=r)
                data.attendance_days.append({"parent_source_label":"Students & Attendance","parent_sheet":sheet,"parent_row":r,"day_of_month":c-9,"raw_status_code":raw_s,"normalized_status_code":norm,"source_label":"Students & Attendance","sheet":sheet,"row":r,"column":get_column_letter(c)})
            if any(not blank(v) for v in vals[4:9]):
                meas_bad=False
                for c,name,allow_zero in ((5,"age_years",True),(6,"height_cm",False),(7,"weight_kg",False),(8,"reported_bmi_value",False)):
                    v=vals[c-1]
                    if blank(v): continue
                    nv=numeric(v)
                    if nv is None:
                        data.failures["Students & Attendance"].append([sheet,r,get_column_letter(c),"measurement",f"Non-numeric {name}",v]); meas_bad=True
                    elif (allow_zero and nv<0) or (not allow_zero and nv<=0):
                        data.failures["Students & Attendance"].append([sheet,r,get_column_letter(c),"measurement",f"{name} violates approved database check constraint",v]); meas_bad=True
                if not meas_bad:
                    data.measurements.append({"parent_source_label":"Students & Attendance","parent_sheet":sheet,"parent_row":r,"age_years":numeric(vals[4]),"height_cm":numeric(vals[5]),"weight_kg":numeric(vals[6]),"reported_bmi_value":numeric(vals[7]),"reported_bmi_category":sval(vals[8]),"source_label":"Students & Attendance","sheet":sheet,"row":r})
    ac=Counter(attendance_keys); attn_dup={k for k,n in ac.items() if n>1}
    for a in data.attendance_months:
        a["duplicate_candidate"]=(a["student_id"],a["academic_year"],a["grade"],a["month_number"]) in attn_dup
    data.summary["Students & Attendance"]={"student_source_rows":student_source_rows,"students_migratable":len(data.students),"student_failed_rows":student_failed_rows,"enrollment_source_cells":enrollment_source_cells,"enrollments_migratable":len(data.enrollments),"enrollment_failures":enrollment_failures,"attendance_source_rows":attendance_source_rows,"attendance_months_migratable":len(data.attendance_months),"attendance_failed_rows":attendance_failed_rows,"attendance_days_migratable":len(data.attendance_days),"measurements_migratable":len(data.measurements),"attendance_duplicate_groups":len(attn_dup),"attendance_duplicate_members":sum(ac[k] for k in attn_dup),"unresolved_grade_labels":dict(unresolved),"attendance_status_normalized":dict(attendance_status_norm),"whitespace_transformations":{f"{a!r}->{b!r}":n for (a,b),n in transforms.items()},"failure_entries":len(data.failures["Students & Attendance"])}

    all_keys=[]; invalid_scores=Counter()
    for w in manifest["workbooks"]:
        if w["kind"]!="assessment": continue
        label=w["label"]; awb=load_workbook(source_dir/w["filename"],data_only=True,read_only=True)
        file_rows=failed_rows=populated=invalid=gs=comps=acts=placeholders=0; file_indices=[]
        for ws in awb.worksheets:
            if ws.sheet_state!="visible" or " - " not in ws.title: continue
            rows=rows_values(ws)
            if len(rows)<3: continue
            grade,subject=[x.strip() for x in ws.title.split(" - ",1)]
            remember(data.subjects,subject,source_label=label,sheet=ws.title,row=1)
            data.grade_subjects.append({"grade":grade,"subject":subject,"source_label":label,"sheet":ws.title,"row":1}); gs+=1
            r2,r3=list(rows[1]),list(rows[2]); maxc=max(len(r2),len(r3))
            avg=next((c for c in range(1,maxc+1) if (sval(r3[c-1] if c-1<len(r3) else None) or "").lower()=="average score"),maxc)
            activity_cols=[]; current_comp=None; comp_order=0; activity_order=0
            for c in range(5,avg):
                comp=sval(r2[c-1] if c-1<len(r2) else None)
                if comp:
                    current_comp=comp; comp_order+=1; activity_order=0; comps+=1
                    data.competencies.append({"grade":grade,"subject":subject,"competency_order":comp_order,"competency_name":comp,"source_label":label,"sheet":ws.title,"row":2})
                act=sval(r3[c-1] if c-1<len(r3) else None)
                if act and act.startswith("<activity_name>"): act="<activity_name>"
                if current_comp is not None: activity_order+=1
                activity_cols.append((c,act,comp_order,activity_order))
                if act and current_comp is not None:
                    acts+=1; placeholders+=int(act=="<activity_name>")
                    data.activities.append({"grade":grade,"subject":subject,"competency_order":comp_order,"activity_order":activity_order,"activity_name":act,"is_placeholder":act=="<activity_name>","source_label":label,"sheet":ws.title,"row":3,"column":get_column_letter(c)})
            for r,vals0 in enumerate(rows[3:],4):
                vals=list(vals0)
                if all(blank(v) for v in vals): continue
                file_rows+=1; sid=sval(vals[0] if len(vals)>0 else None); ay=sval(vals[2] if len(vals)>2 else None); term=sval(vals[3] if len(vals)>3 else None); bad=False
                if not sid:
                    data.failures[label].append([ws.title,r,"A","assessment","Missing required student_id",vals[0] if vals else None]); bad=True
                elif sid not in valid_students:
                    data.failures[label].append([ws.title,r,"A","assessment","student_id not present in valid Students master row",sid]); bad=True
                if not ay:
                    data.failures[label].append([ws.title,r,"C","assessment","Missing required academic_year",vals[2] if len(vals)>2 else None]); bad=True
                if not term:
                    data.failures[label].append([ws.title,r,"D","assessment","Missing required assessment period",vals[3] if len(vals)>3 else None]); bad=True
                elif term not in PERIODS:
                    data.failures[label].append([ws.title,r,"D","assessment","Unrecognized assessment period",term]); bad=True
                if bad: failed_rows+=1; continue
                remember(data.academic_years,ay,source_label=label,sheet=ws.title,row=r); remember(data.periods,term,source_label=label,sheet=ws.title,row=r)
                idx=len(data.assessments); file_indices.append(idx); all_keys.append((sid,ay,grade,subject,term,idx))
                data.assessments.append({"student_id":sid,"academic_year":ay,"grade":grade,"subject":subject,"period":term,"reported_average_score":numeric(vals[avg-1] if avg-1<len(vals) else None),"duplicate_candidate":False,"source_label":label,"sheet":ws.title,"row":r})
                for c,act,co,ao in activity_cols:
                    v=vals[c-1] if c-1<len(vals) else None
                    if blank(v): continue
                    populated+=1; raw=sval(v)
                    if not act:
                        data.failures[label].append([ws.title,r,get_column_letter(c),"assessment_score","Score present under blank activity header",v]); invalid+=1; invalid_scores[raw]+=1; continue
                    parsed=parse_assessment_score(raw)
                    if parsed is None:
                        data.failures[label].append([ws.title,r,get_column_letter(c),"assessment_score","Invalid assessment score; allowed integer 0-10 or A",raw]); invalid+=1; invalid_scores[raw]+=1; continue
                    num,code=parsed
                    if code=="A":
                        data.score_codes.setdefault("A",{"source_label":"LKG","sheet":"Dropdown-Range","row":13})
                    data.scores.append({"assessment_source_label":label,"assessment_sheet":ws.title,"assessment_row":r,"grade":grade,"subject":subject,"competency_order":co,"activity_order":ao,"raw_score_value":raw,"numeric_score":num,"score_code":code,"source_label":label,"sheet":ws.title,"row":r,"column":get_column_letter(c)})
        data.summary[label]={"assessment_source_rows":file_rows,"assessments_migratable":len(file_indices),"assessment_failed_rows":failed_rows,"populated_score_cells":populated,"scores_migratable":sum(1 for s in data.scores if s["source_label"]==label),"invalid_score_cells":invalid,"grade_subjects":gs,"competencies":comps,"activities":acts,"placeholder_activities":placeholders,"failure_entries":len(data.failures[label])}

    kc=Counter((sid,ay,g,s,t) for sid,ay,g,s,t,_ in all_keys); dup={k for k,n in kc.items() if n>1}
    for sid,ay,g,s,t,idx in all_keys: data.assessments[idx]["duplicate_candidate"]=(sid,ay,g,s,t) in dup
    for label in [w["label"] for w in manifest["workbooks"] if w["kind"]=="assessment"]:
        c=Counter((sid,ay,g,s,t) for sid,ay,g,s,t,idx in all_keys if data.assessments[idx]["source_label"]==label)
        data.summary[label]["duplicate_groups"]=sum(1 for n in c.values() if n>1); data.summary[label]["duplicate_members"]=sum(n for n in c.values() if n>1)
    data.summary["_totals"]={"valid_students":len(valid_students),"assessment_duplicate_groups":len(dup),"assessment_duplicate_members":sum(kc[k] for k in dup),"attendance_duplicate_groups":len(attn_dup),"attendance_duplicate_members":sum(ac[k] for k in attn_dup),"invalid_score_values":dict(invalid_scores),"grade_subjects":len(data.grade_subjects),"competencies":len(data.competencies),"activities":len(data.activities),"placeholder_activities":sum(a["is_placeholder"] for a in data.activities),"failure_entries_total":sum(len(v) for v in data.failures.values())}
    return data

def write_failure_log(data:StageEData,output:Path):
    output.parent.mkdir(parents=True,exist_ok=True); wb=Workbook(); wb.remove(wb.active)
    hf=PatternFill("solid",fgColor="1F4E78"); hfont=Font(color="FFFFFF",bold=True); wf=PatternFill("solid",fgColor="FCE4D6")
    s=wb.create_sheet("Summary"); s.append(["Source","Metric","Count / Value"])
    for c in s[1]: c.fill=hf; c.font=hfont
    for src,metrics in data.summary.items():
        for k,v in metrics.items(): s.append([src,k,json.dumps(v,ensure_ascii=False) if isinstance(v,(dict,list)) else v])
    s.freeze_panes="A2"; s.auto_filter.ref=s.dimensions; s.column_dimensions["A"].width=24; s.column_dimensions["B"].width=36; s.column_dimensions["C"].width=75
    for src in FAILURE_SHEET_ORDER:
        ws=wb.create_sheet(src[:31]); ws.append(["Source Sheet","Row","Column","Failure Scope","Reason","Raw Value"])
        for c in ws[1]: c.fill=hf; c.font=hfont
        rows=data.failures.get(src,[])
        if rows:
            for x in rows: ws.append(x)
            for row in ws.iter_rows(min_row=2):
                for c in row: c.fill=wf
        else: ws.append(["","","","","No migration failures identified",""])
        ws.freeze_panes="A2"; ws.auto_filter.ref=ws.dimensions
        for i,w in enumerate([28,10,10,20,62,30],1): ws.column_dimensions[get_column_letter(i)].width=w
        for row in ws.iter_rows():
            for c in row: c.alignment=Alignment(vertical="top",wrap_text=True)
    wb.save(output)

def expected_counts(data):
    return {"source_files":len(data.source_files),"academic_years":len(data.academic_years),"grades":len(data.grades),"subjects":len(data.subjects),"assessment_periods":len(data.periods),"assessment_score_codes":len(data.score_codes),"attendance_status_codes":len(data.attendance_codes),"students":len(data.students),"student_grade_enrollments":len(data.enrollments),"grade_subjects":len(data.grade_subjects),"assessment_competencies":len(data.competencies),"assessment_activities":len(data.activities),"student_assessments":len(data.assessments),"assessment_scores":len(data.scores),"attendance_months":len(data.attendance_months),"attendance_days":len(data.attendance_days),"student_measurements":len(data.measurements)}
